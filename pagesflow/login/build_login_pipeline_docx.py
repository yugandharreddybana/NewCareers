#!/usr/bin/env python3
"""Generate login-pipeline.docx from structured content and diagram PNGs."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
DIAGRAMS = BASE / "diagrams"
OUT = BASE / "login-pipeline.docx"


def add_title(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


def add_diagram(doc: Document, title: str, png_name: str, caption: str | None = None) -> None:
    add_title(doc, title, level=2)
    png = DIAGRAMS / png_name
    if png.exists():
        doc.add_picture(str(png), width=Inches(6.5))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_para(doc, f"[Diagram missing: {png_name}]", bold=True)
    if caption:
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("CareerOps Login Pipeline — Complete Reference", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Sign-in page for returning users. Supports email/password and Google OAuth, "
        "optional Remember Me, and a built-in jumbled word CAPTCHA below the password field "
        "(no Google reCAPTCHA iframe on the email form). Guests only; authenticated users "
        "are redirected into the app. Shows a session-expired banner when redirected from "
        "a silent refresh failure."
    )

    # ── Architecture ──
    add_title(doc, "1. Architecture Overview", 1)
    add_para(
        doc,
        "The login stack spans four layers: React UI (Login.tsx + AuthContext), Axios HTTP client "
        "with CSRF and silent refresh, Express middleware BFF (cookie issuance + validation + HMAC proxy), "
        "and Spring Boot backend (AuthService, WordCaptchaService, TwoFactorService, GoogleOAuthService, JwtService)."
    )
    add_diagram(doc, "1.1 Auth stack layers", "auth-layers.png",
                "Frontend → Middleware BFF → Java backend with HMAC trust boundary")

    # ── Route ──
    add_title(doc, "2. Route and Access Control", 1)
    add_table(doc, ["Property", "Value"], [
        ["URL", "/login"],
        ["Guard", "GuestRoute"],
        ["Layout", "Standalone (LoginPageShell + AuthIntelligencePanel right panel)"],
        ["Redirects (logged-in)", "onboarded → /dashboard or state.from path; not onboarded → /onboarding"],
    ])
    add_title(doc, "2.1 Query parameters", 2)
    add_table(doc, ["Param", "Effect"], [
        ["?reason=session_expired", "Shows banner: Your session expired. Please sign in again."],
        ["?email=", "Prefills email field; stripped from URL via history.replaceState"],
    ])
    add_title(doc, "2.2 GuestRoute redirect logic", 2)
    add_bullets(doc, [
        "While AuthContext.loading is true, shows PageLoader.",
        "If user is already authenticated, Navigate to loginRedirectTarget(user, from).",
        "safeRedirectPath rejects open redirects and auth-loop paths (/login, /signup, etc.).",
        "Default: onboarded users → /dashboard; non-onboarded → /onboarding.",
    ])

    # ── Fields ──
    add_title(doc, "3. Fields and Inputs", 1)
    add_para(doc, "Field order on the form: Email → Password → Security check → Remember me → Submit.")
    add_table(doc, ["Field", "Required", "Validation", "When shown"], [
        ["Email", "Yes", "HTML5 type=email, trimmed on submit", "Always"],
        ["Password", "Yes", "Client min length 8 + server enforcement", "Always"],
        ["Security check", "Yes (prod/staging)", "Jumbled 5 chars; case-insensitive; captchaToken = challengeId:answer", "When LOGIN_WORD_CAPTCHA_REQUIRED"],
        ["Remember me", "No", "Passed to signIn as rememberMe: true", "Always"],
    ])
    add_title(doc, "3.1 Security check (Word CAPTCHA) UX", 2)
    add_bullets(doc, [
        "On mount, WordCaptchaField fetches GET /auth/captcha/challenge.",
        "Server returns challengeId + SVG with 5 jumbled characters (per-letter rotation/color).",
        "Answer is embedded only in SVG — never returned as plain text in JSON.",
        "User types left-to-right; frontend builds captchaToken as challengeId:answer.",
        "Refresh button loads a new challenge. Failed login also refreshes the challenge.",
        "Each challengeId is one-time use (consumed on verify). TTL: 10 minutes server-side.",
        "Google sign-in does NOT use word CAPTCHA (uses reCAPTCHA on consent sheet when CAPTCHA_ENABLED).",
    ])

    # ── Actions ──
    add_title(doc, "4. User Actions", 1)
    add_table(doc, ["Action", "Trigger", "Result"], [
        ["Sign in (email)", "Form submit", "Client validate → AuthContext.signIn → POST /auth/login"],
        ["Sign in (Google)", "LoginGoogleButton credential", "Consent sheet if needed → signInWithGoogle → POST /auth/google"],
        ["Link Google to password account", "409 LINK_REQUIRES_VERIFICATION", "Password + reCAPTCHA → POST /auth/google/link/confirm"],
        ["2FA verify", "6-digit OTP submit", "completeTwoFactor → POST /auth/two-factor/verify"],
        ["Refresh security check", "Button on captcha row", "New GET /auth/captcha/challenge"],
        ["Forgot password", "Link", "Navigate to /forgot-password"],
        ["Sign up", "Footer link", "Navigate to /signup"],
        ["Show/hide password", "Eye toggle", "Toggles input type"],
    ])

    # ── Pipeline phases ──
    add_title(doc, "5. Login Pipeline — Phased Flow", 1)
    add_para(doc, "The login journey is split into seven phases. Each phase has a sequence diagram below.")

    add_title(doc, "5.0 Phase 0 — Route entry and session bootstrap", 2)
    add_bullets(doc, [
        "GuestRoute wraps /login; triggers AuthContext.refresh() on mount.",
        "shouldSkipInitialSessionProbe(): on guest pages with no tokens, skip /auth/me (fast path).",
        "If co_refresh cookie exists, ensureFreshSession → POST /auth/refresh → GET /auth/me.",
        "If user already authenticated, GuestRoute redirects before Login.tsx renders.",
    ])
    add_diagram(doc, "Diagram: Phase 0 — Session bootstrap", "phase0-bootstrap.png")

    add_title(doc, "5.1 Phase 1 — Page mount: word CAPTCHA challenge", 2)
    add_bullets(doc, [
        "WordCaptchaField mounts when LOGIN_WORD_CAPTCHA_REQUIRED (prod/staging) or captchaRequired flag set after failed login.",
        "WordCaptchaService.createChallenge(): 5 chars from A-Z/2-9 (no ambiguous chars), shuffled with rotation in SVG.",
        "Disabled in local dev: auth.login.word-captcha.required=false on backend; LOGIN_WORD_CAPTCHA_REQUIRED=false on frontend.",
    ])
    add_diagram(doc, "Diagram: Phase 1 — CAPTCHA challenge", "phase1-captcha.png")

    add_title(doc, "5.2 Phase 2 — Email/password login (primary path)", 2)
    add_para(doc, "AuthService.login() security behaviors:")
    add_bullets(doc, [
        "Account lockout checked BEFORE password verify (lockedUntil in future → generic 401).",
        "Word CAPTCHA verified when loginWordCaptchaRequired is true.",
        "Google-only accounts (null password hash) → generic 401 with timing-safe dummy BCrypt.",
        "Always runs BCrypt verify (DUMMY_HASH if user not found) — prevents timing/email enumeration.",
        "Failed attempts: atomic incrementFailedAttempts; at 5 failures → lock 15 minutes + ACCOUNT_LOCKED audit.",
        "Success: resetFailedAttempts, assertAccountActive, audit LOGIN.",
        "If 2FA enabled: issueChallengeToken → requiresTwoFactor response (no cookies yet).",
        "If no 2FA: createAuthResponse → JWT access (15min) + refresh token (SHA-256 at rest, UA/subnet binding).",
    ])
    add_diagram(doc, "Diagram: Phase 2 — Email/password login", "phase2-email-login.png",
                "Primary login path with captcha, lockout, 2FA branch, and cookie issuance")

    add_title(doc, "5.3 Phase 3 — Two-factor authentication completion", 2)
    add_bullets(doc, [
        "Login.tsx catches TWO_FACTOR_REQUIRED error with challengeToken; shows OtpInput (6 digits).",
        "Middleware sets co_remember cookie at login step if rememberMe was checked (before 2FA).",
        "POST /auth/two-factor/verify: resolveRememberMe from cookies + body.",
        "TwoFactorService.verifyLoginCode validates TOTP against challenge.",
        "Success: createAuthResponse + issueAuthCookies; AuthContext.setUser.",
    ])
    add_diagram(doc, "Diagram: Phase 3 — 2FA verification", "phase3-2fa.png")

    add_title(doc, "5.4 Phase 4 — Google sign-in", 2)
    add_bullets(doc, [
        "LoginGoogleButton returns Google idToken via GIS credential callback.",
        "resolveGoogleLoginConsent(sessionStorage): may open GoogleConsentSheet for terms/AI/marketing/analytics.",
        "authenticateWithGoogle: verifyIdToken → find by googleSub or email.",
        "Existing googleSub user: assert active, reset failures, update lastLoginAt.",
        "Email exists with password but no googleSub: 409 LINK_REQUIRES_VERIFICATION → Phase 5.",
        "Email exists, linkable (no password): auto-set googleSub.",
        "New user: requires terms + AI consent; createGoogleUser + provisionDefaultOrg.",
        "Google path uses requireRecaptchaWhenConfigured on consent sheet (not word CAPTCHA).",
    ])
    add_diagram(doc, "Diagram: Phase 4 — Google sign-in", "phase4-google.png")

    add_title(doc, "5.5 Phase 5 — Google account linking", 2)
    add_bullets(doc, [
        "Triggered when email/password account exists and user signs in with Google for same email.",
        "writePendingGoogleLink(idToken) to sessionStorage; in-memory googleLinkToken state.",
        "User enters account password + reCAPTCHA (if CAPTCHA_ENABLED).",
        "confirmGoogleLink: verify idToken + BCrypt password → set googleSub, GOOGLE provider.",
        "Page refresh loses link flow — user must click Google again (UI warns).",
    ])
    add_diagram(doc, "Diagram: Phase 5 — Google link confirm", "phase5-google-link.png")

    add_title(doc, "5.6 Phase 6 — Post-login redirect", 2)
    add_bullets(doc, [
        "AuthContext.setUser triggers GuestRoute re-render.",
        "safeRedirectPath(state.from) used if user was sent from protected route.",
        "Otherwise: user.onboarded ? /dashboard : /onboarding.",
        "On success: clearPendingSignup(), clearOnboardingVerification(), syncLocalAnalyticsConsentToBackend().",
    ])
    add_diagram(doc, "Diagram: Phase 6 — Post-login redirect", "phase6-redirect.png")

    # ── Auth and session ──
    add_title(doc, "6. Auth, Cookies, and Tokens", 1)
    add_table(doc, ["Cookie / artifact", "Purpose", "Lifetime"], [
        ["co_session", "Access JWT (HttpOnly)", "15 minutes"],
        ["co_refresh", "Refresh JWT (HttpOnly)", "Session cookie OR 30 days (remember me)"],
        ["co_remember", "Persistent refresh flag", "30 days when remember me checked"],
        ["co_csrf", "CSRF token for mutating requests", "Session"],
        ["tokenStore (memory)", "Access token metadata for Bearer header", "In-memory only"],
    ])
    add_title(doc, "6.1 Remember me behavior", 2)
    add_table(doc, ["Remember me", "Refresh storage", "TTL"], [
        ["Unchecked", "HttpOnly co_refresh session cookie", "Browser session"],
        ["Checked", "HttpOnly co_refresh + co_remember", "30 days"],
    ])
    add_para(doc, "Access token handling is identical in both modes. Refresh token is NEVER stored in sessionStorage.")
    add_title(doc, "6.2 Middleware cookie issuance", 2)
    add_bullets(doc, [
        "authJsonResponse: Set-Cookie co_session + co_refresh; JSON body returns { user } only (LSA-071).",
        "completeAuthResponse in api.ts: if no token in JSON, POST /auth/refresh to get in-memory access.",
        "2FA login step: middleware may set co_remember before returning challengeToken.",
    ])
    add_title(doc, "6.3 Refresh token security (backend)", 2)
    add_bullets(doc, [
        "Raw refresh token never stored in DB — SHA-256 hash only.",
        "Token rotation on refresh; reuse detection revokes entire token family.",
        "Binding hash: User-Agent + /24 subnet; mismatch → 401 + family revoke.",
    ])

    # ── API ──
    add_title(doc, "7. API Endpoints", 1)
    add_table(doc, ["User action", "Frontend", "Middleware (/api/v1)", "Java (/api)"], [
        ["Load captcha", "authApi.getWordCaptchaChallenge", "GET /auth/captcha/challenge", "WordCaptchaService.createChallenge"],
        ["Email login", "authApi.login", "POST /auth/login", "AuthService.login"],
        ["2FA verify", "authApi.verifyTwoFactor", "POST /auth/two-factor/verify", "AuthService.verifyTwoFactorLogin"],
        ["Google login", "authApi.google", "POST /auth/google", "AuthService.authenticateWithGoogle"],
        ["Google link", "authApi.confirmGoogleLink", "POST /auth/google/link/confirm", "AuthService.confirmGoogleLink"],
        ["Session probe", "authApi.me", "GET /auth/me", "AuthService.me"],
        ["Silent refresh", "authApi.refresh / ensureFreshSession", "POST /auth/refresh", "AuthService.refresh"],
        ["Logout", "authApi.logout", "POST /auth/logout", "AuthService.logout"],
    ])
    add_title(doc, "7.1 Login request body (email path)", 2)
    add_para(doc, "{ email, password, rememberMe?, captchaToken? } where captchaToken = challengeId:answer (required when word CAPTCHA enabled).")

    # ── File map ──
    add_title(doc, "8. File Map", 1)
    add_title(doc, "8.1 Frontend", 2)
    add_table(doc, ["Role", "Path"], [
        ["Page", "frontend/src/pages/Login.tsx"],
        ["Word CAPTCHA UI", "frontend/src/components/auth/WordCaptchaField.tsx"],
        ["Shell / intelligence panel", "frontend/src/components/auth/LoginPageShell.tsx, AuthIntelligencePanel.tsx"],
        ["Google button", "frontend/src/components/auth/LoginGoogleButton.tsx"],
        ["Google consent sheet", "frontend/src/components/auth/GoogleConsentSheet.tsx"],
        ["Copy constants", "frontend/src/components/auth/loginCopy.ts"],
        ["Auth state", "frontend/src/context/AuthContext.tsx"],
        ["HTTP client", "frontend/src/services/api.ts"],
        ["Token store", "frontend/src/lib/tokenStore.ts"],
        ["Route guard", "frontend/src/components/ProtectedRoute.tsx (GuestRoute)"],
        ["Pending Google link", "frontend/src/lib/pendingGoogleLink.ts"],
        ["Pending Google consents", "frontend/src/lib/pendingGoogleConsents.ts"],
    ])
    add_title(doc, "8.2 Middleware", 2)
    add_table(doc, ["Role", "Path"], [
        ["Auth routes", "middleware/src/routes/auth.routes.ts"],
        ["Auth guard", "middleware/src/authGuard.ts"],
        ["JWT verification", "middleware/src/jwtVerification.ts"],
        ["Remember me resolver", "middleware/src/authRememberMe.js"],
        ["Rate limiting", "middleware/src/rateLimiter.ts (loginLimiter)"],
    ])
    add_title(doc, "8.3 Backend", 2)
    add_table(doc, ["Role", "Path"], [
        ["Controller", "backend/src/main/java/com/careerops/controller/AuthController.java"],
        ["Login + brute-force", "backend/src/main/java/com/careerops/service/AuthService.java"],
        ["Word CAPTCHA", "backend/src/main/java/com/careerops/service/WordCaptchaService.java"],
        ["2FA", "backend/src/main/java/com/careerops/service/TwoFactorService.java"],
        ["Google OAuth", "backend/src/main/java/com/careerops/service/GoogleOAuthService.java"],
        ["JWT", "backend/src/main/java/com/careerops/security/JwtService.java"],
        ["Google reCAPTCHA (consent/link)", "backend/src/main/java/com/careerops/service/CaptchaService.java"],
    ])

    # ── Edge cases ──
    add_title(doc, "9. Edge Cases and Error Handling", 1)
    add_bullets(doc, [
        "Session expired redirect: Axios interceptor / silent refresh failure → /login?reason=session_expired. First-time guest visits to protected routes go to plain /login (no banner).",
        "2FA remember me: rememberMe passed in POST /auth/two-factor/verify body; co_remember cookie may be set at login step.",
        "Google link refresh: link step-up token in-memory + sessionStorage; refreshing page requires signing in with Google again.",
        "Email prefill: ?email= pre-populates field; query param removed from address bar immediately after read.",
        "One-time CAPTCHA challenge: each challengeId consumed on verify; refresh or retry fetches new one.",
        "Challenge TTL: 10 minutes server-side; expired challenges return invalid captcha.",
        "Google-only account: password login fails with generic error (no enumeration).",
        "Account lockout: 5+ failures → 15-minute lock; still returns generic 401.",
        "No VITE_RECAPTCHA_SITE_KEY needed for email word CAPTCHA (unlike onboarding reCAPTCHA).",
        "Failed login with captchaRequired flag in API error re-enables CAPTCHA UI even in dev.",
        "Middleware loginLimiter applied to POST /auth/login and /auth/two-factor/verify.",
        "CSRF: mutating requests include X-CSRF-Token from co_csrf cookie.",
        "HMAC: middleware → Java requests signed with APP_INTERNAL_SECRET (see shared/auth-infrastructure.md).",
        "Generic error messages: all login failures return same message to prevent email enumeration.",
        "Timing attack mitigation: BCrypt always runs even when user not found.",
    ])

    # ── Related docs ──
    add_title(doc, "10. Related Documentation", 1)
    add_bullets(doc, [
        "pagesflow/login/PAGE.md — concise page reference",
        "pagesflow/shared/auth-infrastructure.md — cookies, refresh, axios, HMAC",
        "pagesflow/shared/route-guards.md — GuestRoute, ProtectedRoute, redirect tree",
        "pagesflow/mandatory-fields.md — env vars for auth in dev/prod",
        "pagesflow/signup/PAGE.md — deferred signup flow",
        "pagesflow/onboarding/PAGE.md — post-login onboarding for new users",
        "pagesflow/forgot-password/PAGE.md — password reset flow",
    ])

    # ── Env ──
    add_title(doc, "11. Environment Variables (Login-relevant)", 1)
    add_table(doc, ["Variable", "Layer", "Purpose"], [
        ["LOGIN_WORD_CAPTCHA_REQUIRED / auth.login.word-captcha.required", "Frontend + Java", "Enable word CAPTCHA on email login"],
        ["VITE_RECAPTCHA_SITE_KEY / RECAPTCHA_SECRET_KEY", "Frontend + Java", "Google consent sheet and link form reCAPTCHA"],
        ["JWT_PUBLIC_KEY / JWT keys", "Middleware + Java", "Access token issue and verify"],
        ["APP_INTERNAL_SECRET", "Middleware + Java", "HMAC signing for BFF → backend"],
        ["COOKIE_NAME, COOKIE_SAMESITE", "Middleware", "Session cookie configuration"],
        ["TRUSTED_PROXY", "Middleware + Java", "Client IP for rate limits and refresh binding"],
    ])

    doc.add_paragraph()
    footer = doc.add_paragraph("Generated for CareerOps pagesflow/login. Diagram sources: pagesflow/login/diagrams/*.mmd")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in footer.runs:
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
